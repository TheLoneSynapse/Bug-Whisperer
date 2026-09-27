"""TEST: openpyxl — ValueError from an invalid cell reference."""

from openpyxl import Workbook

wb = Workbook()
ws = wb.active

# WRONG (ValueError: ZZZZ99999 is not a valid coordinate or range)
ws["ZZZZ99999"] = "Hello"

# CORRECT
# ws["A1"] = "Hello"
