"""DEMO: openpyxl mistake — using an invalid cell reference.

Run with:
    .\run.bat demos\demo_openpyxl.py
"""

from openpyxl import Workbook

wb = Workbook()
ws = wb.active

# WRONG (ValueError: ZZZZ99999 is not a valid coordinate or range)
ws["ZZZZ99999"] = "Hello"

# CORRECT (use a valid cell like A1, B2, etc.)
ws["A1"] = "Hello"
print("Cell A1 value:", ws["A1"].value)
